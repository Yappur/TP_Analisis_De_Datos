"""Genera un Excel por ejercicio en la carpeta excel/ (excel/ejercicio_01.xlsx, ...).

Todos los ejercicios:      python exportar_excel.py
Solo algunos:              python exportar_excel.py 1 5 10
Un ejercicio desde su      python -m ejercicios.ej05      (ejecutar desde la carpeta del proyecto)
propio archivo:
"""
import importlib
import sys
from math import ceil
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.properties import PageSetupProperties

BASE = Path(__file__).parent
sys.path.insert(0, str(BASE))  # para poder importar "ejercicios" y "estadistica" desde cualquier lado
CARPETA_SALIDA = BASE / "excel"

COLOR_TITULO = "166A7A"
COLOR_ENCABEZADO = "D9F3F7"
COLOR_NOTA = "FFF7F3"
LADO = Side(style="thin", color="B8CED4")
BORDES = Border(left=LADO, right=LADO, top=LADO, bottom=LADO)

ANCHOS = {"A": 34, "B": 18, "C": 18, "D": 14, "E": 14, "F": 4}  # el gráfico va desde la columna G
CARACTERES_ANCHO_TEXTO = 95   # lo que entra en una línea de las celdas combinadas A:E
FILAS_GRAFICO = 19            # filas que ocupa un gráfico


def escribir_fila(hoja, fila, valores, encabezado=False):
    for columna, valor in enumerate(valores, start=1):
        celda = hoja.cell(row=fila, column=columna, value=valor)
        celda.border = BORDES
        es_numero = isinstance(valor, (int, float))
        celda.alignment = Alignment(horizontal="right" if es_numero else "left",
                                    vertical="center", wrap_text=True)
        if encabezado:
            celda.font = Font(bold=True, color="12333D")
            celda.fill = PatternFill("solid", fgColor=COLOR_ENCABEZADO)
            celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def escribir_texto(hoja, fila, texto, relleno=None, negrita=False):
    """Texto largo: celdas A:E combinadas, con la altura de fila ajustada."""
    hoja.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=5)
    celda = hoja.cell(row=fila, column=1, value=texto)
    celda.alignment = Alignment(wrap_text=True, vertical="top")
    celda.font = Font(bold=negrita)
    if relleno:
        celda.fill = PatternFill("solid", fgColor=relleno)
    hoja.row_dimensions[fila].height = 16 * max(1, ceil(len(texto) / CARACTERES_ANCHO_TEXTO)) + 2


def agregar_grafico(hoja, grafico, fila_encabezado, fila_ultima, fila_ancla, titulo):
    series = len(grafico["series"])
    g = BarChart()
    g.type = "col"
    g.title = titulo
    g.y_axis.title = "Frecuencia"
    g.gapWidth = 60
    g.width, g.height = 17, 9
    # Sin esto, las versiones nuevas de Excel pueden ocultar los ejes
    g.x_axis.delete = False
    g.y_axis.delete = False
    if series == 1:
        g.legend = None
    else:
        g.legend.position = "b"
    datos = Reference(hoja, min_col=2, max_col=1 + series, min_row=fila_encabezado, max_row=fila_ultima)
    categorias = Reference(hoja, min_col=1, min_row=fila_encabezado + 1, max_row=fila_ultima)
    g.add_data(datos, titles_from_data=True)
    g.set_categories(categorias)
    hoja.add_chart(g, f"G{fila_ancla}")


def escribir_bloque(hoja, fila, bloque):
    fila_inicio = fila
    hoja.cell(row=fila, column=1, value=bloque["subtitulo"]).font = Font(bold=True, size=13, color=COLOR_TITULO)
    fila += 1

    escribir_fila(hoja, fila, bloque["columnas"], encabezado=True)
    fila += 1
    for datos in bloque["filas"]:
        escribir_fila(hoja, fila, datos)
        fila += 1

    for nota in bloque["notas"]:
        escribir_texto(hoja, fila, f"Nota: {nota}", relleno=COLOR_NOTA)
        fila += 1

    grafico = bloque.get("grafico")
    if grafico:
        # Tabla auxiliar con los datos del gráfico (el gráfico de Excel necesita celdas)
        fila += 1
        hoja.cell(row=fila, column=1, value="Datos del gráfico").font = Font(bold=True, color=COLOR_TITULO)
        fila += 1
        fila_encabezado = fila
        escribir_fila(hoja, fila, ["Etiqueta"] + [s["nombre"] for s in grafico["series"]], encabezado=True)
        fila += 1
        for i, etiqueta in enumerate(grafico["etiquetas"]):
            escribir_fila(hoja, fila, [etiqueta] + [s["valores"][i] for s in grafico["series"]])
            fila += 1
        agregar_grafico(hoja, grafico, fila_encabezado, fila - 1, fila_inicio, bloque["subtitulo"])
        fila = max(fila, fila_inicio + FILAS_GRAFICO)  # que el próximo bloque no pise el gráfico

    return fila + 2


def configurar_hoja(hoja):
    for letra, ancho in ANCHOS.items():
        hoja.column_dimensions[letra].width = ancho
    hoja.sheet_view.showGridLines = False
    # Para imprimir / pasar a PDF: horizontal y entra todo el ancho en una página
    hoja.page_setup.orientation = "landscape"
    hoja.page_setup.fitToWidth = 1
    hoja.page_setup.fitToHeight = 0
    hoja.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)


def exportar(modulo):
    """Crea excel/ejercicio_NN.xlsx a partir de un módulo de ejercicios/ (TITULO, ENUNCIADO, resolver)."""
    numero = int(Path(modulo.__file__).stem[2:])  # ej05.py -> 5
    libro = Workbook()
    hoja = libro.active
    hoja.title = f"Ejercicio {numero}"
    configurar_hoja(hoja)

    hoja["A1"] = modulo.TITULO
    hoja["A1"].font = Font(bold=True, size=16, color=COLOR_TITULO)
    escribir_texto(hoja, 2, modulo.ENUNCIADO)

    fila = 4
    for bloque in modulo.resolver():
        fila = escribir_bloque(hoja, fila, bloque)

    CARPETA_SALIDA.mkdir(parents=True, exist_ok=True)
    ruta = CARPETA_SALIDA / f"ejercicio_{numero:02d}.xlsx"
    try:
        libro.save(ruta)
    except PermissionError:
        raise SystemExit(f"No se pudo guardar {ruta}: cerrá el archivo en Excel y volvé a ejecutar.")
    return ruta


def numeros_disponibles():
    return sorted(int(p.stem[2:]) for p in (BASE / "ejercicios").glob("ej[0-9][0-9].py"))


if __name__ == "__main__":
    pedidos = [int(a) for a in sys.argv[1:]] or numeros_disponibles()
    for n in pedidos:
        try:
            modulo = importlib.import_module(f"ejercicios.ej{n:02d}")
        except ModuleNotFoundError:
            print(f"Ejercicio {n}: no existe ejercicios/ej{n:02d}.py")
            continue
        print(f"Ejercicio {n}: {exportar(modulo)}")
