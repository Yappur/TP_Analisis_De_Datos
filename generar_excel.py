"""Genera excel/ejercicios.xlsx a partir de los ejercicios del TP."""
from pathlib import Path
import importlib

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


BASE = Path(__file__).parent
CARPETA_EJERCICIOS = BASE / "ejercicios"
SALIDA = BASE / "excel" / "ejercicios.xlsx"
COLOR_TITULO = "166A7A"
COLOR_ENCABEZADO = "D9F3F7"
COLOR_NOTA = "FFF7F3"
BORDES = Border(
    left=Side(style="thin", color="D7E4E8"),
    right=Side(style="thin", color="D7E4E8"),
    top=Side(style="thin", color="D7E4E8"),
    bottom=Side(style="thin", color="D7E4E8"),
)


def escribir_fila(hoja, fila, valores, negrita=False, relleno=None):
    for columna, valor in enumerate(valores, start=1):
        celda = hoja.cell(row=fila, column=columna, value=valor)
        celda.alignment = Alignment(vertical="top", wrap_text=True)
        celda.border = BORDES
        if negrita:
            celda.font = Font(bold=True, color="12333D")
        if relleno:
            celda.fill = PatternFill("solid", fgColor=relleno)


def ajustar_columnas(hoja):
    for columna in hoja.columns:
        ancho = 12
        letra = get_column_letter(columna[0].column)
        for celda in columna:
            if celda.value is not None:
                ancho = max(ancho, min(len(str(celda.value)) + 2, 48))
        hoja.column_dimensions[letra].width = ancho


def escribir_bloque(hoja, fila, bloque):
    hoja.cell(row=fila, column=1, value=bloque["subtitulo"])
    hoja.cell(row=fila, column=1).font = Font(bold=True, size=13, color=COLOR_TITULO)
    fila += 1

    escribir_fila(hoja, fila, bloque["columnas"], negrita=True, relleno=COLOR_ENCABEZADO)
    fila += 1
    for datos in bloque["filas"]:
        escribir_fila(hoja, fila, datos)
        fila += 1

    for nota in bloque["notas"]:
        hoja.cell(row=fila, column=1, value=f"Nota: {nota}")
        hoja.cell(row=fila, column=1).fill = PatternFill("solid", fgColor=COLOR_NOTA)
        hoja.cell(row=fila, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        fila += 1

    grafico = bloque.get("grafico")
    if grafico:
        fila += 1
        hoja.cell(row=fila, column=1, value="Datos del gráfico")
        hoja.cell(row=fila, column=1).font = Font(bold=True, color=COLOR_TITULO)
        fila += 1
        encabezados = ["Etiqueta"] + [serie["nombre"] for serie in grafico["series"]]
        escribir_fila(hoja, fila, encabezados, negrita=True, relleno=COLOR_ENCABEZADO)
        fila += 1
        for indice, etiqueta in enumerate(grafico["etiquetas"]):
            valores = [etiqueta] + [serie["valores"][indice] for serie in grafico["series"]]
            escribir_fila(hoja, fila, valores)
            fila += 1

    return fila + 2


def cargar_ejercicios():
    archivos = sorted(CARPETA_EJERCICIOS.glob("ej[0-9][0-9].py"))
    for archivo in archivos:
        numero = int(archivo.stem[2:])
        modulo = importlib.import_module(f"ejercicios.{archivo.stem}")
        yield numero, modulo


def generar_excel():
    SALIDA.parent.mkdir(exist_ok=True)
    libro = Workbook()
    resumen = libro.active
    resumen.title = "Resumen"
    escribir_fila(resumen, 1, ["Ejercicio", "Título"], negrita=True, relleno=COLOR_ENCABEZADO)

    for numero, modulo in cargar_ejercicios():
        escribir_fila(resumen, numero + 1, [numero, modulo.TITULO])
        hoja = libro.create_sheet(f"Ejercicio {numero}")
        hoja["A1"] = modulo.TITULO
        hoja["A1"].font = Font(bold=True, size=16, color=COLOR_TITULO)
        hoja["A2"] = modulo.ENUNCIADO
        hoja["A2"].alignment = Alignment(wrap_text=True, vertical="top")

        fila = 4
        for bloque in modulo.resolver():
            fila = escribir_bloque(hoja, fila, bloque)

        hoja.freeze_panes = "A4"
        ajustar_columnas(hoja)

    ajustar_columnas(resumen)
    libro.save(SALIDA)
    return SALIDA


if __name__ == "__main__":
    ruta = generar_excel()
    print(f"Excel generado en {ruta}")
